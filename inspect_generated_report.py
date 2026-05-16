from pathlib import Path
from docx import Document

p = Path(r"C:\Users\Zz\Desktop\生产实习\output\report_final.docx")
d = Document(p)
print("paragraphs", len(d.paragraphs), "tables", len(d.tables), "inline_shapes", len(d.inline_shapes), "sections", len(d.sections))
for ti, t in enumerate(d.tables[:2]):
    print("TABLE", ti)
    for row in t.rows:
        print([c.text for c in row.cells])
    print("---")
for s in ["一、选题介绍", "二、需求分析", "三、数据库概念结构设计", "四、数据库逻辑结构设计", "五、数据库的实施", "六、应用系统开发", "七、问题记录和解决方案", "八、总结"]:
    print(s, any(s in p.text for p in d.paragraphs))
