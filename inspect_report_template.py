from pathlib import Path
from docx import Document

path = Path(r"C:\Users\Zz\Desktop\生产实习\output\专业认知实习报告-黄钰宗-2415304249.docx")
doc = Document(path)
print("paragraphs", len(doc.paragraphs), "tables", len(doc.tables), "sections", len(doc.sections))
for i, para in enumerate(doc.paragraphs[:160]):
    text = para.text.strip().replace("\n", " ")
    if text:
        print(f"P{i}: {text}")
for ti, table in enumerate(doc.tables):
    print(f"TABLE {ti}: {len(table.rows)}x{len(table.columns)}")
    for ri, row in enumerate(table.rows[:12]):
        vals = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
        print("  R", ri, vals)
