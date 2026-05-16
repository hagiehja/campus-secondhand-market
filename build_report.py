from __future__ import annotations

import math
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"C:\Users\Zz\Desktop\生产实习")
OUT = ROOT / "output"
TEMPLATE = OUT / "专业认知实习报告-黄钰宗-2415304249.docx"
FINAL_ASCII = OUT / "report_final.docx"
FINAL_CN = OUT / "专业认知实习报告-24软工2班-黄钰宗-2415304249.docx"
LOGIN_SHOT = OUT / "login-page-reference-style.png"
ER_IMAGE = OUT / "campus_market_er_chen.png"
ARCH_IMAGE = OUT / "campus_market_architecture.png"
MODULE_IMAGE = OUT / "campus_market_modules.png"
HOME_IMAGE = OUT / "campus_market_home_mock.png"


def font_path(name: str) -> str:
    candidates = [
        Path(r"C:\Windows\Fonts") / name,
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simsun.ttc"),
    ]
    for p in candidates:
        if p.exists():
            return str(p)
    return ""


FONT_REG = font_path("msyh.ttc")
FONT_BOLD = font_path("msyhbd.ttc")


def f(size: int, bold: bool = False):
    path = FONT_BOLD if bold and FONT_BOLD else FONT_REG
    return ImageFont.truetype(path, size) if path else ImageFont.load_default()


def set_cell_text(cell, text: str, bold: bool = False, size: int = 11):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(size)
    r.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="B7C6BE", sz="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), sz)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def delete_paragraph(paragraph):
    p = paragraph._element
    parent = p.getparent()
    parent.remove(p)
    paragraph._p = paragraph._element = None


def style_document(doc: Document):
    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.4)
        section.right_margin = Cm(2.4)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    for name, size in [("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)]:
        try:
            style = styles[name]
        except KeyError:
            continue
        style.font.name = "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0x0F, 0x56, 0x4E)
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(6)


def add_run_cn(paragraph, text: str, bold=False, size=None, color=None):
    run = paragraph.add_run(text)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run


def add_body_paragraph(doc, text: str = "", bold_prefix: str | None = None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        add_run_cn(p, bold_prefix, bold=True)
        add_run_cn(p, text[len(bold_prefix):])
    else:
        add_run_cn(p, text)
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.5
    return p


def add_bullet(doc, text: str):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.line_spacing = 1.3
    add_run_cn(p, text)


def set_table_style(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def draw_rounded(draw, box, radius, fill, outline=None, width=1):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def center_text(draw, box, text, font, fill=(20, 35, 31)):
    x1, y1, x2, y2 = box
    bbox = draw.multiline_textbbox((0, 0), text, font=font, spacing=4, align="center")
    w = bbox[2] - bbox[0]
    h = bbox[3] - bbox[1]
    draw.multiline_text((x1 + (x2 - x1 - w) / 2, y1 + (y2 - y1 - h) / 2), text, font=font, fill=fill, spacing=4, align="center")


def make_diamond(cx, cy, w, h):
    return [(cx, cy - h // 2), (cx + w // 2, cy), (cx, cy + h // 2), (cx - w // 2, cy)]


def draw_entity(draw, box, name):
    draw.rectangle(box, fill="#F8FFFC", outline="#0F766E", width=3)
    center_text(draw, box, name, f(25, True), fill=(8, 47, 41))


def draw_attr(draw, cx, cy, text, key=False):
    w, h = 180, 58
    box = (cx - w // 2, cy - h // 2, cx + w // 2, cy + h // 2)
    draw.ellipse(box, fill="#FFFFFF", outline="#364943", width=2)
    center_text(draw, box, text, f(20, key), fill=(15, 50, 45))
    if key:
        tw = draw.textbbox((0, 0), text, font=f(20, True))[2]
        draw.line((cx - tw / 2, cy + 14, cx + tw / 2, cy + 14), fill="#0F766E", width=2)


def draw_relation(draw, cx, cy, text):
    pts = make_diamond(cx, cy, 112, 72)
    draw.polygon(pts, fill="#ECFDF5", outline="#0F766E")
    draw.line([pts[0], pts[1], pts[2], pts[3], pts[0]], fill="#0F766E", width=3)
    center_text(draw, (cx - 56, cy - 36, cx + 56, cy + 36), text, f(20, True), fill=(7, 65, 56))


def draw_line(draw, a, b, label=None, offset=(0, 0)):
    draw.line((a[0], a[1], b[0], b[1]), fill="#42534C", width=3)
    if label:
        x = (a[0] + b[0]) / 2 + offset[0]
        y = (a[1] + b[1]) / 2 + offset[1]
        draw.text((x, y), label, fill="#0F766E", font=f(22, True))


def create_er_diagram():
    img = Image.new("RGB", (1900, 1300), "#FFFFFF")
    d = ImageDraw.Draw(img)
    d.text((70, 42), "校园二手交易系统 E-R 图", font=f(34, True), fill="#0F564E")

    entities = {
        "用户": (120, 560, 310, 640),
        "商品": (640, 560, 830, 640),
        "分类": (640, 210, 830, 290),
        "订单": (1160, 560, 1350, 640),
        "商品图片": (640, 910, 860, 990),
        "收藏": (1160, 210, 1350, 290),
        "消息": (1550, 560, 1740, 640),
        "评价": (1160, 910, 1350, 990),
    }
    for name, box in entities.items():
        draw_entity(d, box, name)

    # attributes, placed away from relationship lines.
    attrs = [
        (215, 430, "用户编号", True, (215, 560)),
        (85, 730, "用户名", False, (165, 640)),
        (300, 760, "角色", False, (260, 640)),
        (735, 450, "商品编号", True, (735, 560)),
        (560, 710, "标题", False, (675, 640)),
        (800, 725, "价格", False, (780, 640)),
        (900, 545, "成色", False, (830, 595)),
        (735, 95, "分类编号", True, (735, 210)),
        (920, 245, "分类名称", False, (830, 250)),
        (1255, 430, "订单编号", True, (1255, 560)),
        (1445, 555, "金额", False, (1350, 590)),
        (1235, 735, "状态", False, (1235, 640)),
        (750, 1130, "图片URL", False, (750, 990)),
        (1255, 1130, "评分", False, (1255, 990)),
        (1645, 430, "消息编号", True, (1645, 560)),
        (1790, 720, "内容", False, (1700, 640)),
    ]
    for cx, cy, text, key, target in attrs:
        draw_attr(d, cx, cy, text, key)
        d.line((cx, cy, target[0], target[1]), fill="#67756F", width=2)

    # relations
    draw_relation(d, 475, 600, "发布")
    draw_line(d, (310, 600), (419, 600), "1", (-75, -38))
    draw_line(d, (531, 600), (640, 600), "n", (30, -38))

    draw_relation(d, 735, 415, "属于")
    draw_line(d, (735, 290), (735, 379), "1", (16, -34))
    draw_line(d, (735, 451), (735, 560), "n", (16, 18))

    draw_relation(d, 995, 600, "下单")
    draw_line(d, (830, 600), (939, 600), "1", (-72, -38))
    draw_line(d, (1051, 600), (1160, 600), "n", (28, -38))

    draw_relation(d, 735, 780, "拥有")
    draw_line(d, (735, 640), (735, 744), "1", (16, -30))
    draw_line(d, (735, 816), (735, 910), "n", (16, 14))

    draw_relation(d, 995, 250, "收藏")
    draw_line(d, (830, 250), (939, 250), "n", (-72, -38))
    draw_line(d, (1051, 250), (1160, 250), "n", (28, -38))
    draw_line(d, (215, 560), (215, 250), None)
    draw_line(d, (215, 250), (939, 250), "1", (-315, -38))

    draw_relation(d, 1455, 600, "咨询")
    draw_line(d, (1350, 600), (1399, 600), "1", (-36, -38))
    draw_line(d, (1511, 600), (1550, 600), "n", (10, -38))

    draw_relation(d, 1255, 780, "评价")
    draw_line(d, (1255, 640), (1255, 744), "1", (16, -30))
    draw_line(d, (1255, 816), (1255, 910), "n", (16, 14))

    d.text((70, 1225), "说明：下划线属性为主键；1、n 表示一对多或多对多基数。", font=f(22), fill="#33443D")
    img.save(ER_IMAGE, quality=95)


def create_architecture_diagram():
    img = Image.new("RGB", (1600, 820), "#F6FAF7")
    d = ImageDraw.Draw(img)
    d.text((70, 45), "系统技术架构图", font=f(38, True), fill="#0F564E")
    boxes = [
        ((80, 190, 410, 390), "Vue 前端", "HTML / CSS / JS\n商品浏览、登录、订单、支付入口"),
        ((620, 190, 980, 390), "Spring Boot 后端", "REST API\n业务校验、订单状态、支付回调"),
        ((1180, 120, 1500, 300), "MySQL 5.7", "192.168.24.129:3307\n用户、商品、订单、评价数据"),
        ((1180, 390, 1500, 570), "Redis", "192.168.24.129:6379\n登录状态、热点缓存、浏览计数"),
    ]
    for box, title, desc in boxes:
        draw_rounded(d, box, 24, "#FFFFFF", "#0F766E", 3)
        center_text(d, (box[0], box[1] + 22, box[2], box[1] + 78), title, f(30, True), fill=(9, 73, 64))
        center_text(d, (box[0] + 20, box[1] + 90, box[2] - 20, box[3] - 18), desc, f(23), fill=(48, 66, 60))
    d.line((410, 290, 620, 290), fill="#0F766E", width=5)
    d.polygon([(620, 290), (595, 275), (595, 305)], fill="#0F766E")
    d.text((455, 242), "HTTP/JSON", font=f(24, True), fill="#0F766E")
    d.line((980, 260, 1180, 210), fill="#0F766E", width=5)
    d.polygon([(1180, 210), (1152, 203), (1162, 232)], fill="#0F766E")
    d.line((980, 320, 1180, 480), fill="#0F766E", width=5)
    d.polygon([(1180, 480), (1152, 476), (1168, 454)], fill="#0F766E")
    d.text((1015, 180), "JDBC", font=f(24, True), fill="#0F766E")
    d.text((1015, 420), "Spring Data Redis", font=f(24, True), fill="#0F766E")
    draw_rounded(d, (220, 610, 1380, 720), 20, "#E8F7EF", "#C6EBDD", 2)
    center_text(d, (240, 620, 1360, 710), "开发与运行：Java 17 + Spring Boot 3.2.4 + Vue 3 + Vite；数据库和 Redis 均部署在 192.168.24.129 虚拟机。", f(24), fill=(22, 74, 65))
    img.save(ARCH_IMAGE, quality=95)


def create_modules_diagram():
    img = Image.new("RGB", (1600, 900), "#FFFFFF")
    d = ImageDraw.Draw(img)
    d.text((70, 45), "系统功能模块图", font=f(38, True), fill="#0F564E")
    center = (640, 410, 960, 530)
    draw_rounded(d, center, 20, "#0F766E", "#0F766E", 3)
    center_text(d, center, "校园二手交易系统", f(32, True), fill=(255, 255, 255))
    modules = [
        ((120, 160, 400, 280), "用户登录", "账号密码登录、角色识别、登录状态"),
        ((120, 540, 400, 660), "商品管理", "商品列表、分类筛选、发布演示"),
        ((660, 90, 940, 210), "订单管理", "创建订单、查询订单、状态追踪"),
        ((1120, 160, 1400, 280), "支付管理", "模拟支付、支付宝接口预留"),
        ((1120, 540, 1400, 660), "消息沟通", "咨询消息、交易提醒、安全提示"),
        ((660, 670, 940, 790), "数据统计", "市场价值、分类数量、交易数据"),
    ]
    for box, title, desc in modules:
        draw_rounded(d, box, 18, "#F2FBF6", "#0F766E", 3)
        center_text(d, (box[0], box[1] + 12, box[2], box[1] + 58), title, f(27, True), fill=(9, 73, 64))
        center_text(d, (box[0] + 18, box[1] + 62, box[2] - 18, box[3] - 12), desc, f(21), fill=(49, 67, 60))
        d.line(((box[0]+box[2])//2, (box[1]+box[3])//2, 800, 470), fill="#7AAEA1", width=3)
    img.save(MODULE_IMAGE, quality=95)


def create_home_mock():
    img = Image.new("RGB", (1600, 1000), "#F4F7F2")
    d = ImageDraw.Draw(img)
    draw_rounded(d, (30, 30, 270, 970), 12, "#FFFFFF", "#DCE7DF", 2)
    d.text((70, 80), "校园集市", font=f(34, True), fill="#0F564E")
    for i, nav in enumerate(["首页", "分类", "发布商品", "我的订单", "消息"]):
        y = 150 + i * 70
        fill = "#DFF7ED" if i == 0 else "#FFFFFF"
        draw_rounded(d, (60, y, 240, y + 48), 10, fill, "#E1EAE5", 1)
        d.text((95, y + 12), nav, font=f(22, True), fill="#0F564E")
    draw_rounded(d, (310, 30, 1250, 250), 12, "#0F766E", "#0F766E", 2)
    d.text((350, 70), "发现同学正在出手的好物", font=f(38, True), fill="#FFFFFF")
    d.text((350, 130), "分类、搜索、价格排序与商品浏览均已接入 Spring Boot 后端接口。", font=f(24), fill="#DDFCF2")
    chips = ["全部", "最新", "价格低到高", "成色良好"]
    for i, chip in enumerate(chips):
        x = 330 + i * 190
        draw_rounded(d, (x, 285, x + 150, 335), 25, "#E8F7EF" if i == 0 else "#FFFFFF", "#0F766E", 2)
        center_text(d, (x, 285, x + 150, 335), chip, f(22, True), fill=(15, 86, 78))
    products = [
        ("高等数学同济第七版", "教材资料", "¥32", "#E8F7EF"),
        ("计算机网络教材", "教材资料", "¥30", "#EDF3FF"),
        ("Java程序设计实训书", "教材资料", "¥28", "#FFF7E6"),
        ("蓝牙耳机", "数码产品", "¥118", "#F8FAFC"),
        ("校园通勤自行车", "运动器材", "¥260", "#EAFBF5"),
        ("机械键盘青轴", "数码产品", "¥139", "#F3F6FA"),
    ]
    for i, (title, cat, price, fill) in enumerate(products):
        row, col = divmod(i, 3)
        x = 310 + col * 310
        y = 365 + row * 270
        draw_rounded(d, (x, y, x + 280, y + 235), 12, "#FFFFFF", "#E0E7E2", 2)
        draw_rounded(d, (x, y, x + 280, y + 105), 12, fill, fill, 1)
        d.ellipse((x + 190, y + 20, x + 245, y + 75), fill="#FFFFFF")
        d.text((x + 25, y + 128), cat, font=f(20, True), fill="#0F766E")
        d.text((x + 25, y + 160), title, font=f(22, True), fill="#0F201A")
        d.text((x + 25, y + 205), price, font=f(25, True), fill="#EF5A2F")
    draw_rounded(d, (1280, 30, 1565, 970), 12, "#FFFFFF", "#DCE7DF", 2)
    d.text((1320, 80), "热门推荐", font=f(30, True), fill="#0F201A")
    for i, text in enumerate(["教材资料成交活跃", "数码产品浏览上升", "宿舍用品需求稳定"]):
        d.text((1320, 145 + i * 78), text, font=f(22, True), fill="#0F766E")
        d.text((1320, 177 + i * 78), "后端实时返回统计数据", font=f(18), fill="#5C6C66")
    draw_rounded(d, (1320, 520, 1535, 700), 12, "#17251F", "#17251F", 2)
    d.text((1350, 565), "今日市场估值", font=f(20), fill="#DFF7ED")
    d.text((1350, 610), "¥ 9658", font=f(42, True), fill="#FB923C")
    img.save(HOME_IMAGE, quality=95)


def add_figure(doc, image_path: Path, caption: str, width_cm: float = 15.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run_cn(cap, caption, size=9, color=(92, 106, 99))
    cap.paragraph_format.space_after = Pt(8)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_style(table)
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, size=10)
        shade_cell(table.rows[0].cells[i], "E8F7EF")
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], str(val), size=9)
    if widths:
        for row in table.rows:
            for i, cell in enumerate(row.cells):
                cell.width = Cm(widths[i])
    doc.add_paragraph()
    return table


def build_doc():
    create_er_diagram()
    create_architecture_diagram()
    create_modules_diagram()
    create_home_mock()

    doc = Document(TEMPLATE)
    style_document(doc)
    # Fill cover table.
    for table in doc.tables:
        for row in table.rows:
            key = row.cells[0].text.strip()
            if "学" in key and "号" in key:
                set_cell_text(row.cells[1], "2415304249")
            if "姓" in key and "名" in key:
                set_cell_text(row.cells[1], "黄钰宗")
            if "班" in key and "级" in key:
                set_cell_text(row.cells[1], "24软工2班")
    # Remove template instructions after date line.
    for p in list(doc.paragraphs[11:]):
        delete_paragraph(p)

    doc.add_page_break()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run_cn(title, "校园二手交易系统设计与实现", bold=True, size=18, color=(15, 86, 78))
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run_cn(subtitle, "专业认知实习报告", size=12, color=(80, 96, 90))

    add_table(doc, ["项目", "内容"], [
        ["学生信息", "黄钰宗，学号 2415304249，24软工2班"],
        ["项目名称", "校园二手交易系统"],
        ["技术栈", "Java、Spring Boot、MySQL 5.7、Redis、Vue、HTML、CSS、JavaScript"],
        ["中间件部署", "MySQL：192.168.24.129:3307；Redis：192.168.24.129:6379"],
        ["本地运行", "后端：http://127.0.0.1:8080；前端：http://127.0.0.1:5173/?real=1"],
    ], widths=[3.2, 12.0])

    doc.add_heading("一、选题介绍", level=1)
    add_body_paragraph(doc, "本次专业认知实习选择“校园二手交易系统”作为课程设计主题。校园内教材、数码产品、生活用品、运动器材等闲置物品流通需求较高，但传统微信群或线下张贴信息存在商品信息分散、交易记录不完整、沟通效率低和安全性不足等问题。系统以校内学生为主要用户，围绕商品展示、分类检索、登录认证、下单支付、订单查询和消息提醒等流程进行设计，实现一个面向校园场景的轻量级二手交易平台。")
    add_body_paragraph(doc, "项目的目的在于把数据库课程中的概念结构设计、逻辑结构设计、SQL 实施、触发器、视图、索引以及前后端开发流程串联起来。通过该系统，可以训练从需求分析到数据库建模、从 SQL 脚本到 Spring Boot 接口、从 Vue 页面到真实数据展示的完整开发能力。项目的意义不仅在于完成课程要求，也在于形成一个贴近学生生活的应用原型，为后续扩展实名认证、聊天、支付和后台管理等功能奠定基础。")

    doc.add_heading("二、需求分析", level=1)
    add_body_paragraph(doc, "系统面向三类角色：管理员、卖家和买家。管理员负责维护用户、商品分类、商品和交易数据；卖家负责发布商品、维护商品图片、处理订单和回复买家咨询；买家负责浏览商品、按分类筛选、收藏商品、咨询卖家、下单购买并完成支付。")
    add_figure(doc, MODULE_IMAGE, "图 1 系统功能模块图", 15.2)
    add_body_paragraph(doc, "核心业务流程为：用户进入系统后进行登录；首页从后端接口读取在售商品和分类统计；买家选择商品后创建订单；订单创建后商品被锁定，避免重复交易；买家可选择模拟支付或进入支付宝支付页面；支付完成后订单状态更新，后续可扩展评价与消息通知。")
    add_table(doc, ["角色", "主要需求", "对应功能"], [
        ["买家", "快速找到校内二手商品并完成下单", "商品浏览、分类筛选、搜索、下单、支付、订单查询"],
        ["卖家", "发布并管理自己的闲置商品", "商品发布、商品状态维护、订单处理、消息回复"],
        ["管理员", "维护平台基础数据和交易秩序", "用户管理、分类管理、商品审核、日志查看"],
    ], widths=[2.4, 6.4, 6.4])

    doc.add_heading("三、数据库概念结构设计", level=1)
    add_body_paragraph(doc, "根据需求分析，系统抽象出用户、商品、分类、商品图片、交易订单、收藏、消息、评价和操作日志等实体。用户与商品之间存在“发布”关系，分类与商品之间存在“属于”关系，商品与订单之间存在“下单/生成”关系，商品与图片之间存在“拥有”关系，用户与商品之间通过收藏表形成多对多关系，订单与评价之间形成一对多关系。")
    add_figure(doc, ER_IMAGE, "图 2 校园二手交易系统 E-R 图", 16.0)

    doc.add_heading("四、数据库逻辑结构设计", level=1)
    add_body_paragraph(doc, "将 E-R 图转换为关系模型时，每个实体转换为一张数据表，一对多联系通过在多端加入外键实现，多对多联系通过中间表实现。例如用户和商品的一对多关系通过 product.seller_id 外键关联 sys_user.id；用户收藏商品的多对多关系通过 favorite(user_id, product_id) 中间表实现。")
    add_table(doc, ["表名", "主键", "主要字段", "说明"], [
        ["sys_user", "id", "username, password_hash, real_name, student_no, role, status", "系统用户表，保存买家、卖家和管理员信息"],
        ["category", "id", "name, parent_id, sort_no, status", "商品分类表，支持父子分类"],
        ["product", "id", "seller_id, category_id, title, price, condition_level, status", "商品表，保存商品基本信息和交易状态"],
        ["product_image", "id", "product_id, image_url, sort_no", "商品图片表，一件商品可对应多张图片"],
        ["trade_order", "id", "order_no, product_id, buyer_id, seller_id, amount, status", "交易订单表，记录下单、支付和完成状态"],
        ["favorite", "id", "user_id, product_id", "收藏表，维护用户与商品的多对多关系"],
        ["message", "id", "product_id, sender_id, receiver_id, content, is_read", "消息表，保存买卖双方咨询内容"],
        ["review", "id", "order_id, reviewer_id, target_user_id, score, content", "评价表，保存交易完成后的评分与评价"],
        ["operation_log", "id", "biz_type, biz_id, action, detail", "操作日志表，记录重要业务操作"],
    ], widths=[2.3, 1.7, 7.1, 4.4])

    doc.add_heading("五、数据库的实施", level=1)
    add_body_paragraph(doc, "数据库使用 MySQL 5.7，部署在虚拟机 192.168.24.129 的 3307 端口。系统 SQL 文件包括基础建表脚本、演示数据脚本、扩展商品数据脚本和支付扩展脚本。核心建表语句如下所示：")
    add_table(doc, ["对象", "SQL 示例"], [
        ["数据库", "CREATE DATABASE campus_secondhand_market DEFAULT CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci;"],
        ["用户表", "CREATE TABLE sys_user (... username VARCHAR(50) UNIQUE, role ENUM('ADMIN','SELLER','BUYER'), status TINYINT ...);"],
        ["商品表", "CREATE TABLE product (... seller_id BIGINT, category_id BIGINT, price DECIMAL(10,2), status ENUM('ON_SALE','LOCKED','SOLD') ...);"],
        ["订单表", "CREATE TABLE trade_order (... order_no VARCHAR(32) UNIQUE, buyer_id BIGINT, seller_id BIGINT, status ENUM('PENDING','PAID','FINISHED','CANCELLED') ...);"],
    ], widths=[2.4, 12.8])
    add_body_paragraph(doc, "为了满足课程设计中视图和触发器要求，数据库中设计了在售商品视图 v_on_sale_product，并设计了订单触发器用于校验订单与锁定商品。")
    add_table(doc, ["数据库对象", "作用"], [
        ["视图 v_on_sale_product", "整合 product、category、sys_user 和 product_image，前端首页直接查询在售商品展示所需字段。"],
        ["索引 idx_product_search", "提升商品标题、分类和状态筛选场景下的查询效率。"],
        ["触发器 trg_order_before_insert_validate_order", "下单前校验商品是否仍处于可售状态。"],
        ["触发器 trg_order_after_insert_lock_product", "订单创建后自动将商品状态更新为 LOCKED，避免重复下单。"],
        ["触发器 trg_order_after_update_finish_product", "订单完成后自动将商品状态更新为 SOLD。"],
    ], widths=[4.4, 10.8])
    add_body_paragraph(doc, "系统导入了 43 条在售商品演示数据，覆盖教材资料、数码产品、生活用品、运动器材、服饰鞋包五类商品。接口验证时 /api/products?size=1 返回 totalElements=43，说明数据库数据能够被后端正确读取。")

    doc.add_heading("六、应用系统开发", level=1)
    add_body_paragraph(doc, "开发环境为 Java 17、Spring Boot 3.2.4、MySQL 5.7、Redis、Vue 3 和 Vite。后端负责提供 REST API，前端负责页面展示与交互，MySQL 保存持久化业务数据，Redis 预留用于登录状态、商品缓存和热点浏览计数。开发过程中主要使用 Codex 辅助完成数据库脚本、接口联调、前端页面设计和运行验证。")
    add_figure(doc, ARCH_IMAGE, "图 3 系统技术架构图", 15.2)
    add_body_paragraph(doc, "登录模块：前端登录页使用深色校园夜景风格，包含账号、密码、记住我、忘记密码、第三方登录入口等界面元素。用户输入 buyer01 / 123456 后，请求后端 /api/auth/login 接口，后端从 sys_user 表中查询用户并返回用户编号、角色和 token。")
    if LOGIN_SHOT.exists():
        add_figure(doc, LOGIN_SHOT, "图 4 登录界面截图", 13.4)
    add_body_paragraph(doc, "商品浏览模块：前端首页调用 /api/products、/api/categories 和 /api/market/summary 接口，展示商品列表、分类数量和市场统计信息。页面支持按分类切换、关键词搜索、价格排序和商品卡片浏览。")
    add_figure(doc, HOME_IMAGE, "图 5 首页商品展示界面示意", 15.2)
    add_table(doc, ["功能", "接口/SQL", "处理说明"], [
        ["登录", "POST /api/auth/login；SELECT * FROM sys_user WHERE username=?", "验证账号密码，返回当前用户信息。"],
        ["商品查询", "GET /api/products；SELECT * FROM v_on_sale_product WHERE status='ON_SALE'", "查询在售商品，前端以卡片形式展示。"],
        ["分类统计", "GET /api/categories；COUNT(product.id) GROUP BY category", "返回各分类商品数量，用于侧边栏和筛选。"],
        ["创建订单", "POST /api/orders；INSERT INTO trade_order(...)", "生成订单号，并通过触发器锁定商品。"],
        ["模拟支付", "POST /api/orders/{id}/pay/mock；UPDATE trade_order SET status='PAID'", "课程演示中用于模拟真实支付成功。"],
        ["订单查询", "GET /api/orders?buyerId=3；SELECT * FROM trade_order WHERE buyer_id=?", "展示当前买家的订单列表和支付状态。"],
    ], widths=[2.4, 5.8, 7.0])

    doc.add_heading("七、问题记录和解决方案", level=1)
    add_table(doc, ["问题", "原因分析", "解决方案"], [
        ["MySQL 使用 3307 端口", "数据库运行在虚拟机中，不是本机默认 3306。", "在 application.yml 中配置 jdbc:mysql://192.168.24.129:3307/campus_secondhand_market。"],
        ["前端页面功能无法切换", "导航最初是静态链接，没有绑定 Vue 状态。", "新增 activeView、navItems 和 switchView 方法，实现首页、分类、发布、订单、消息切换。"],
        ["商品数量太少", "初始演示数据不足，页面不符合真实平台效果。", "补充 campus_secondhand_market_more_products.sql，将在售商品扩展到 43 条。"],
        ["登录页视觉不够贴近参考图", "原页面留白较多，登录卡与背景风格不统一。", "重构登录页为夜景校园背景、悬浮商品卡和深色玻璃登录卡。"],
        ["真实支付宝支付接入复杂", "真实支付需要开放平台应用、公钥私钥和公网回调地址。", "保留支付宝接口位置，并提供 mockPay 模拟支付完成课程演示。"],
    ], widths=[3.4, 5.5, 6.3])

    doc.add_heading("八、总结", level=1)
    add_body_paragraph(doc, "本次专业认知实习围绕校园二手交易系统完成了从数据库设计到应用系统开发的完整实践。项目的亮点在于没有只停留在静态页面或简单 SQL，而是把 MySQL 5.7、Spring Boot、Redis 和 Vue 前端连接起来，形成了可以真实运行的业务链路。数据库方面，系统设计了用户、商品、分类、订单、收藏、消息、评价等表，并使用视图、索引和触发器增强数据查询和业务约束；后端方面，通过 REST API 提供商品列表、分类统计、登录、订单和支付接口；前端方面，完成了校园集市风格的登录页和商品浏览界面，并能与后端数据联动。")
    add_body_paragraph(doc, "在开发过程中，我对数据库概念结构设计和逻辑结构转换有了更直观的认识。以前理解 E-R 图时更多停留在理论层面，这次通过把“用户发布商品”“商品生成订单”“订单产生评价”等关系落到真实表结构中，能明显感受到主键、外键、中间表和触发器在系统中的作用。同时，前后端联调也暴露出一些实际问题，例如端口配置、跨环境数据库连接、演示数据不足和页面交互状态管理等。通过逐步排查和验证，这些问题都得到了处理。")
    add_body_paragraph(doc, "目前系统仍有可以继续完善的地方，例如真实支付宝支付需要进一步申请和配置开放平台参数，商品发布图片上传、后台审核、即时聊天和权限控制也可以继续增强。总体来看，本次实习让我把数据库课程设计要求与真实 Web 项目开发结合起来，进一步理解了数据库并不是孤立存在的脚本文件，而是支撑业务流程、接口设计和用户界面的核心基础。")

    doc.save(FINAL_ASCII)
    shutil.copyfile(FINAL_ASCII, FINAL_CN)
    print(FINAL_ASCII)
    print(FINAL_CN)


if __name__ == "__main__":
    build_doc()
